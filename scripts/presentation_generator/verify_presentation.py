import pptx
import os

path = r'course_materials/semester_4/lectures/lecture_02_thermodynamics_master_layouts.pptx'
print(f"File exists: {os.path.exists(path)}")
print(f"File size: {os.path.getsize(path):,} bytes")

prs = pptx.Presentation(path)
total = len(prs.slides)
print(f"Total slides in presentation: {total}")

print("\n--- Content Lecture Slides (from Slide 30) ---")
for idx in range(29, total):
    slide = prs.slides[idx]
    title = slide.shapes.title.text.strip() if slide.shapes.title and slide.shapes.title.text else "[Без заголовка]"
    # print details
    shapes_info = []
    for s in slide.shapes:
        shapes_info.append(f"{s.name} (w={s.width.pt:.1f}, h={s.height.pt:.1f})")
    print(f"Slide {idx+1}: '{title}' | {len(slide.shapes)} shapes: {', '.join(shapes_info)}")
