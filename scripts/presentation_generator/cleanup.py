from pptx import Presentation
from pptx.util import Inches

def insert_images(file_path):
    prs = Presentation(file_path)
    
    # 2. Clean up shapes on content slides (1 to 6)
    for i in range(1, 7):
        slide = prs.slides[i]
        shapes_to_delete = []
        for shape in slide.shapes:
            # Keep placeholders (Title and Text)
            if shape.is_placeholder:
                continue
            shapes_to_delete.append(shape)
            
        for shape in shapes_to_delete:
            elem = shape._element
            elem.getparent().remove(elem)
            
    # 3. Insert images
    img_map = {
        3: 'phase_water.png',
        4: 'phase_eutectic.png',
        5: 'chart_nucleation.png',
        6: 'chart_gibbs_thomson.png'
    }
    
    left = Inches(6.5)
    top = Inches(1.5)
    width = Inches(5.8)
    
    for slide_idx, img_file in img_map.items():
        slide = prs.slides[slide_idx]
        slide.shapes.add_picture(img_file, left, top, width=width)
        
    prs.save(file_path)
    print("Images inserted.")

if __name__ == "__main__":
    insert_images('course_materials/semester_4/lectures/lecture_02_thermodynamics.pptx')
