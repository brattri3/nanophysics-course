import os
from pptx import Presentation
from pptx.util import Inches

def populate_extended(file_path):
    prs = Presentation(file_path)
    
    slides_data = [
        {'index': 0, 'title': 'Термодинамика наноструктур', 'text': 'Лекция №2 (Расширенная)\nОт основ до наномира\nПреподаватель: Попов Д.А.', 'image': None},
        {'index': 2, 'title': 'Уровень 1. Энергия, Теплота и Работа', 'text': '• Внутренняя энергия U — сумма кинетической энергии молекул и потенциальной энергии их взаимодействия.\n• Теплота Q и Работа A — способы передачи энергии.\n• В школе мы изучали Q = cmΔT и A = FΔx.\n• В термодинамике они неразрывно связаны.', 'image': None},
        {'index': 3, 'title': 'Начала термодинамики и Энтропия', 'text': '• Закон сохранения энергии (I начало): dQ = dU + dA\n• Почему тепло не идет от холодного к горячему?\n• Энтропия S — мера хаоса в системе.\n• Все системы стремятся увеличить свой хаос (II начало).', 'image': 'chart_entropy.png'},
        {'index': 4, 'title': 'Уровень 2. Энергия Гиббса', 'text': '• Как предсказать, пойдет ли реакция, не считая энтропию всей Вселенной?\n• G = H - TS (Энергия Гиббса)\n• Это баланс между стремлением к минимуму энергии (энтальпия H) и максимуму хаоса (энтропия S).\n• Процесс идет самопроизвольно, если ΔG < 0.', 'image': None},
        {'index': 5, 'title': 'Химический потенциал и Равновесие', 'text': '• Химический потенциал μ — это порция энергии Гиббса, приходящаяся на 1 моль вещества.\n• Почему лед плавится ровно при 0 °C?\n• При 0 °C химические потенциалы льда и жидкой воды равны: μ_solid = μ_liquid.\n• В равновесии система не хочет менять фазу.', 'image': None},
        {'index': 6, 'title': 'Уровень 3. Фазовые равновесия', 'text': '• Правило фаз Гиббса: f = k - φ + 2\n• k — число компонентов, φ — число фаз.\n• P-T диаграмма воды: линии показывают условия, где две фазы существуют вместе.\n• Тройная точка (0.01 °C) — сосуществуют сразу 3 фазы (f = 0).', 'image': 'phase_water.png'},
        {'index': 7, 'title': 'Бинарные системы и Эвтектика', 'text': '• Смесь двух компонентов (например, металлов). k = 2.\n• Эвтектика — точка, где жидкость кристаллизуется сразу в две твердые фазы при минимальной температуре.\n• Правило рычага позволяет по графику определить долю каждой фазы в смеси.', 'image': 'phase_eutectic.png'},
        {'index': 8, 'title': 'Уровень 4. Наномир и Поверхностная энергия', 'text': '• Что происходит на границе раздела фаз?\n• Атомы на поверхности имеют оборванные связи (им не хватает соседей).\n• Это создает избыточную энергию — поверхностное натяжение σ.\n• Для наночастиц доля поверхностных атомов огромна!', 'image': 'chart_surface_energy.png'},
        {'index': 9, 'title': 'Зародышеобразование (Барьер)', 'text': '• При кристаллизации объем новой фазы дает выигрыш энергии (ΔG_v < 0).\n• Но создание новой поверхности требует затрат энергии (σ > 0).\n• Возникает барьер. Только частицы крупнее критического радиуса r* могут стабильно расти.', 'image': 'chart_nucleation.png'},
        {'index': 10, 'title': 'Размерный эффект плавления', 'text': '• Эффект Гиббса-Томсона.\n• Из-за огромного искривления поверхности наночастицы испытывают колоссальное внутреннее давление (давление Лапласа).\n• Это буквально "выдавливает" атомы из решетки при нагреве.', 'image': None},
        {'index': 11, 'title': 'Связь с практикой (Синтез в лаборатории)', 'text': '• График показывает, что наночастицы золота (< 5 нм) плавятся при 380 °C вместо 1064 °C!\n• Эти термодинамические эффекты — фундамент ваших лабораторных работ №1-3.\n• Мы управляем барьером зародышеобразования, чтобы получить частицы нужного размера.', 'image': 'chart_gibbs_thomson.png'}
    ]
    
    for data in slides_data:
        slide = prs.slides[data['index']]
        
        if data['index'] > 0:
            shapes_to_delete = []
            text_frames = []
            
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text_frames.append(shape)
                elif not shape.is_placeholder:
                    shapes_to_delete.append(shape)
                    
            for shape in shapes_to_delete:
                sp = shape._element
                sp.getparent().remove(sp)
                
            # Usually the first text frame is title, second is body
            if len(text_frames) > 0: text_frames[0].text = data['title']
            if len(text_frames) > 1: text_frames[1].text = data['text']
                
            if data['image'] and os.path.exists(data['image']):
                slide.shapes.add_picture(data['image'], Inches(5.5), Inches(1.5), width=Inches(6.0))
        else:
            text_frames = [s for s in slide.shapes if s.has_text_frame]
            if len(text_frames) > 0: text_frames[0].text = data['title']
            if len(text_frames) > 1: text_frames[1].text = data['text']
            if len(text_frames) > 2: text_frames[2].text = "2026"
    
    keep_indices = [data['index'] for data in slides_data]
    xml_slides = prs.slides._sldIdLst
    ns = xml_slides.nsmap
    
    for i in range(len(prs.slides) - 1, -1, -1):
        if i not in keep_indices:
            slide_id = prs.slides[i].slide_id
            elem = xml_slides.find(f'.//p:sldId[@id="{slide_id}"]', namespaces=ns)
            if elem is not None:
                xml_slides.remove(elem)
                
    prs.save(file_path)
    print("Extended presentation populated properly.")

if __name__ == "__main__":
    populate_extended('course_materials/semester_4/lectures/lecture_02_thermodynamics_extended.pptx')
