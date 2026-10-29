"""
Автоматизированный верификационный скрипт для Раздела 4 (Слайды 21–26)
и полной презентации Лекции №3:
- Проверка клиренса заголовков >= 15 px (целевой >= 25 px) на слайдах 22–26 над градиентной линией (y=153 px).
- Проверка строго 0 знаков вопроса (?) во всех 26 слайдах, таблицах, заметках докладчика и speech.md.
- Проверка строго 0 запрещенных терминов (лабораторн, экзамен, билет, зачет, тест).
- Проверка Single-Point Occurrence (SPO) автора и кафедры только на Слайде 1.
- Проверка наличия всех 4 компонентов заметок докладчика на слайдах 21–26.
- Проверка полноширинной верстки 12.133" на Слайдах 24 (модульные карточки) и 26 (матричная таблица).
"""

import os
import sys
import pptx
from PIL import Image

PPTX_PATH = "course_materials/semester_4/lectures/lecture_03_colloids.pptx"
SPEECH_PATH = "course_materials/semester_4/lectures/lecture_03_speech.md"
PREVIEWS_DIR = "previews_l03"

FORBIDDEN_TERMS = ["лабораторн", "экзамен", "билет", "зачет", "тест"]

def run_verification():
    print("=" * 78)
    print("АВТОМАТИЗИРОВАННАЯ ВЕРИФИКАЦИЯ РАЗДЕЛА 4 (СЛАЙДЫ 21–26) И ЛЕКЦИИ №3")
    print("=" * 78)

    errors = []

    # 1. Проверка наличия и структуры PPTX
    if not os.path.exists(PPTX_PATH):
        errors.append(f"PPTX файл не найден: {PPTX_PATH}")
        return errors

    prs = pptx.Presentation(PPTX_PATH)
    num_slides = len(prs.slides)
    print(f"[CHECK 1] Количество слайдов в презентации: {num_slides}")
    if num_slides != 26:
        errors.append(f"Количество слайдов {num_slides} != 26")
    else:
        print("  -> PASS: Ровно 26 слайдов.")

    # 2. Клиренс заголовков на слайдах 22–26 (по растровым превью 1920x1080)
    print("\n[CHECK 2] Клиренс заголовков над градиентной линией (y = 153 px) на Слайдах 22–26:")
    grad_line_y = 153
    for s_idx in range(21, 26):  # 0-indexed: 21 -> Slide 22, 25 -> Slide 26
        slide_num = s_idx + 1
        preview_file = os.path.join(PREVIEWS_DIR, f"slide_{slide_num:02d}_preview.png")
        if not os.path.exists(preview_file):
            errors.append(f"Превью не найдено: {preview_file}")
            continue

        with Image.open(preview_file) as img:
            rgb_img = img.convert("RGB")
            w, h = rgb_img.size
            # Сканируем полосу шапки: x от 86 до 1600 px (исключая логотип справа), y от 40 до 151 px
            # Ищем самый нижний пиксель текста (темнее фона #FFFFFF / #F8FAFC, например R < 200)
            lowest_text_y = None
            for y in range(151, 40, -1):
                found = False
                for x in range(86, 1600):
                    r, g, b = rgb_img.getpixel((x, y))
                    # Текст темный: r < 200, g < 200, b < 200
                    if r < 200 and g < 200 and b < 200:
                        lowest_text_y = y
                        found = True
                        break
                if found:
                    break

        if lowest_text_y is not None:
            clearance = grad_line_y - lowest_text_y
            print(f"  Слайд {slide_num}: низ текста y = {lowest_text_y} px, клиренс = {clearance} px (норма >= 15 px, целевой >= 25 px)")
            if clearance < 15:
                errors.append(f"Слайд {slide_num}: недостаточный клиренс заголовка {clearance} px < 15 px")
            elif clearance < 25:
                print(f"    [WARN] Слайд {slide_num}: клиренс {clearance} px в диапазоне 15-25 px")
            else:
                print(f"    -> PASS: Отличный клиренс ({clearance} px >= 25 px)")
        else:
            print(f"  Слайд {slide_num}: текст в зоне шапки не обнаружен (фон чистый)")

    # 3. Проверка 0 знаков вопроса во всей презентации и speech.md
    print("\n[CHECK 3] Проверка отсутствия знаков вопроса (?):")
    total_q_pptx = 0
    for idx, slide in enumerate(prs.slides):
        slide_num = idx + 1
        slide_q = 0
        for s in slide.shapes:
            if s.has_text_frame:
                slide_q += s.text_frame.text.count("?")
            if s.has_table:
                for row in s.table.rows:
                    for cell in row.cells:
                        slide_q += cell.text_frame.text.count("?")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            slide_q += slide.notes_slide.notes_text_frame.text.count("?")
        if slide_q > 0:
            errors.append(f"Слайд {slide_num}: обнаружено {slide_q} знаков вопроса (?)")
            total_q_pptx += slide_q

    print(f"  PPTX знаков вопроса: {total_q_pptx}")

    with open(SPEECH_PATH, "r", encoding="utf-8") as f:
        speech_text = f.read()
    speech_q = speech_text.count("?")
    print(f"  speech.md знаков вопроса: {speech_q}")
    if speech_q > 0:
        errors.append(f"speech.md: обнаружено {speech_q} знаков вопроса (?)")

    if total_q_pptx == 0 and speech_q == 0:
        print("  -> PASS: Строго 0 знаков вопроса во всей презентации и конспекте.")

    # 4. Проверка запрещенных терминов
    print("\n[CHECK 4] Проверка отсутствия запрещенных терминов:")
    found_forbidden = False
    for word in FORBIDDEN_TERMS:
        pptx_matches = 0
        for idx, slide in enumerate(prs.slides):
            slide_text = ""
            for s in slide.shapes:
                if s.has_text_frame:
                    slide_text += " " + s.text_frame.text
                if s.has_table:
                    for row in s.table.rows:
                        for cell in row.cells:
                            slide_text += " " + cell.text_frame.text
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                slide_text += " " + slide.notes_slide.notes_text_frame.text
            cnt = slide_text.lower().count(word)
            if cnt > 0:
                errors.append(f"Слайд {idx+1}: запрещенный термин '{word}' (встретился {cnt} раз)")
                pptx_matches += cnt

        speech_cnt = speech_text.lower().count(word)
        if speech_cnt > 0:
            errors.append(f"speech.md: запрещенный термин '{word}' (встретился {speech_cnt} раз)")

        print(f"  Термин '{word}': в PPTX={pptx_matches}, в speech.md={speech_cnt}")
        if pptx_matches > 0 or speech_cnt > 0:
            found_forbidden = True

    if not found_forbidden:
        print("  -> PASS: 0 запрещенных терминов.")

    # 5. Проверка Single-Point Occurrence (SPO) автора
    print("\n[CHECK 5] Single-Point Occurrence (SPO) автора и кафедры:")
    s1_text = ""
    for s in prs.slides[0].shapes:
        if s.has_text_frame:
            s1_text += " " + s.text_frame.text
    has_author_s1 = "Попов" in s1_text and "Кафедра №3" in s1_text
    print(f"  Слайд 1 содержит реквизиты автора и кафедры: {has_author_s1}")
    if not has_author_s1:
        errors.append("Слайд 1 не содержит полного набора реквизитов автора / кафедры")

    leaks_author = []
    for idx in range(1, len(prs.slides)):
        s_num = idx + 1
        slide_text = ""
        for s in prs.slides[idx].shapes:
            if s.has_text_frame:
                slide_text += " " + s.text_frame.text
            if s.has_table:
                for row in s.table.rows:
                    for cell in row.cells:
                        slide_text += " " + cell.text_frame.text
        if "Попов" in slide_text or "Кафедра №3" in slide_text:
            leaks_author.append(s_num)

    if leaks_author:
        errors.append(f"Нарушение SPO: реквизиты автора найдены на слайдах {leaks_author}")
    else:
        print("  -> PASS: Реквизиты автора и кафедры присутствуют исключительно на Слайде 1.")

    # 6. Проверка 4-компонентных заметок докладчика на слайдах 21–26
    print("\n[CHECK 6] Проверка 4-компонентных заметок докладчика на Слайдах 21–26:")
    req_components = [
        "1. Физический тезис",
        "2. Визуальные акценты",
        "3. Инженерный контекст",
        "4. Смысловой переход"
    ]
    for s_idx in range(20, 26):
        slide_num = s_idx + 1
        slide = prs.slides[s_idx]
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text

        missing = [c for c in req_components if c not in notes]
        if missing:
            errors.append(f"Слайд {slide_num}: в заметках отсутствуют компоненты: {missing}")
            print(f"  Слайд {slide_num}: FAIL - пропущены {missing}")
        else:
            word_count = len(notes.split())
            print(f"  Слайд {slide_num}: PASS - все 4 компонента присутствуют ({word_count} слов)")

    # 7. Проверка модульной полноширинной верстки (12.133") на Слайдах 24 и 26
    print("\n[CHECK 7] Проверка ширины контента на Слайдах 24 и 26 (12.133\"):")
    # Слайд 24: карточки
    s24 = prs.slides[23]
    s24_cards = [s for s in s24.shapes if s.has_text_frame and s.width.inches > 3.0 and s.top.inches > 1.20 and s.top.inches < 2.0]
    if len(s24_cards) >= 3:
        min_left = min(s.left.inches for s in s24_cards)
        max_right = max((s.left.inches + s.width.inches) for s in s24_cards)
        total_span = max_right - min_left
        print(f"  Слайд 24 (карточки): min_left={min_left:.3f}\", max_right={max_right:.3f}\", span={total_span:.3f}\" (цель 12.133\")")
        if abs(total_span - 12.133) > 0.05:
            errors.append(f"Слайд 24: ширина модульной сетки {total_span:.3f}\" != 12.133\"")
        else:
            print("  -> PASS: Слайд 24 имеет полноширинную верстку 12.133\".")
    else:
        errors.append(f"Слайд 24: найдено только {len(s24_cards)} модульных карточек вместо 3")

    # Слайд 26: таблица
    s26 = prs.slides[25]
    tables = [s for s in s26.shapes if s.has_table]
    if tables:
        t = tables[0]
        t_left = t.left.inches
        t_width = t.width.inches
        print(f"  Слайд 26 (таблица): left={t_left:.3f}\", width={t_width:.3f}\" (цель 12.133\")")
        if abs(t_width - 12.133) > 0.05:
            errors.append(f"Слайд 26: ширина таблицы {t_width:.3f}\" != 12.133\"")
        else:
            print("  -> PASS: Слайд 26 имеет полноширинную таблицу 12.133\".")
    else:
        errors.append("Слайд 26: таблица не найдена")

    # 8. Сводный вердикт
    print("\n" + "=" * 78)
    if errors:
        print(f"[FAIL] ОБНАРУЖЕНО {len(errors)} ОШИБОК:")
        for e in errors:
            print(f"  - {e}")
        return False
    else:
        print("[SUCCESS] ВСЕ ВЕРИФИКАЦИОННЫЕ ТЕСТЫ РАЗДЕЛА 4 УСПЕШНО ПРОЙДЕНЫ (PASS)!")
        print("=" * 78)
        return True

if __name__ == "__main__":
    success = run_verification()
    sys.exit(0 if success else 1)
