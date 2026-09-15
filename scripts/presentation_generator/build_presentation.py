import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

OUTPUT_FILE = 'lecture_02_thermodynamics.pptx'
TEMPLATE_FILE = '../../course_materials/templates/guap_16-9_2023.pptx'

if os.path.exists(TEMPLATE_FILE):
    prs = Presentation(TEMPLATE_FILE)
else:
    print(f"Внимание: шаблон {TEMPLATE_FILE} не найден. Используется стандартный.")
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)

def add_title_slide(title, subtitle):
    slide_layout = prs.slide_layouts[0] 
    slide = prs.slides.add_slide(slide_layout)
    title_box = slide.shapes.title
    subtitle_box = slide.placeholders[1]
    
    title_box.text = title
    subtitle_box.text = subtitle

def add_lecture_slide(title, bullet_points, image_path=None):
    slide_layout = prs.slide_layouts[5] # Title only
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.title
    title_box.text = title
    title_box.text_frame.paragraphs[0].font.name = 'Arial'
    title_box.text_frame.paragraphs[0].font.size = Pt(36)
    title_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 64, 128) # #004080
    
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(7.0) if image_path else Inches(15.0)
    height = Inches(7.0)
    
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    for point in bullet_points:
        p = tf.add_paragraph()
        p.text = point
        p.font.name = 'Arial'
        p.font.size = Pt(18)
        p.level = 0
        
    if image_path and os.path.exists(image_path):
        img_left = Inches(8.0)
        img_top = Inches(1.5)
        slide.shapes.add_picture(image_path, img_left, img_top, width=Inches(7.5))

add_title_slide(
    "Термодинамика наноструктур",
    "Лекция №2\nКурс: Физические основы нанотехнологий\nПреподаватель: Попов Д.А."
)

add_lecture_slide(
    'Фундамент термодинамики',
    [
        'Функция состояния U против функций процесса Q и A.',
        'Потенциалы U -> H -> F -> G.',
        'Доминирование изобарно-изотермических условий (T, P = const).'
    ]
)

add_lecture_slide(
    'Химический потенциал μ',
    [
        'μ_i = (∂G / ∂n_i)_T,P.',
        'Термодинамический критерий равновесия фаз: равенство потенциалов.',
        'Если μ_1 > μ_2, масса перемещается во вторую фазу.'
    ]
)

add_lecture_slide(
    'Правило фаз Гиббса и P-T диаграммы',
    [
        'f = k - φ + 2.',
        'Фазовые поля (phi = 1, f = 2): свободное варьирование параметров T и P.',
        'Линии переходов (phi = 2, f = 1): моновариантное равновесие (кипение, плавление).',
        'Тройная точка (phi = 3, f = 0): инвариантное состояние с жестко фиксированными T и P.'
    ],
    'phase_water.png'
)

add_lecture_slide(
    'Бинарные диаграммы и правило рычага',
    [
        'Координаты T — концентрация (x). Линии ликвидуса (расплав) и солидуса (кристалл).',
        'Эвтектика: совместная кристаллизация двух фаз при постоянной T_e.',
        'Правило рычага (коннода): доли фаз обратно пропорциональны отрезкам: m_alpha / m_beta = l_beta / l_alpha.'
    ],
    'phase_eutectic.png'
)

add_lecture_slide(
    'Зародышеобразование: критический радиус r*',
    [
        'Энергия образования зародыша: Delta_G(r) = (4/3)pi r^3 Delta_G_V + 4pi r^2 sigma (Delta_G_V < 0, sigma > 0).',
        'Конкуренция объемного выигрыша (~ -r^3) и поверхностного проигрыша (~ +r^2).',
        'Условие максимума d(Delta_G)/dr = 0 дает критический радиус: r* = -2sigma / Delta_G_V.',
        'Энергетический барьер: зародыши r < r* растворяются; при r > r* растут.'
    ],
    'chart_nucleation.png'
)

add_lecture_slide(
    'Размерная зависимость плавления (Гиббс–Томсон)',
    [
        'Формула Гиббса–Томсона: T_m(r) = T_m,inf * [ 1 - 2sigma_sl / (rho_s * L * r) ].',
        'Физика эффекта: избыточная поверхностная энергия ненасыщенных связей кластера.',
        'Пример золота (Au): спад точки плавления с 1064°C до ~380°C при радиусе r approx 1.5 нм.',
        'Связь с ЛР №1–3: управление дисперсностью наночастиц серебра и золота.'
    ],
    'chart_gibbs_thomson.png'
)

prs.save(OUTPUT_FILE)
print(f'Презентация создана: {OUTPUT_FILE}')
