from pptx import Presentation

def delete_slides(file_path):
    prs = Presentation(file_path)
    
    # We want to KEEP indices 1, 3, 4, 5, 6, 7, 8
    # Which corresponds to:
    # 1: Our Title (Slide 2)
    # 3: Our Block 0 (Slide 4)
    # 4: Our Block 1 (Slide 5)
    # 5: Our Block 2/3 (Slide 6)
    # 6: Our Block 4 (Slide 7)
    # 7: Our Block 5 (Slide 8)
    # 8: Our Block 5/6 (Slide 9)
    keep_indices = [1, 3, 4, 5, 6, 7, 8]
    
    xml_slides = prs.slides._sldIdLst
    ns = xml_slides.nsmap
    
    for i in range(len(prs.slides) - 1, -1, -1):
        if i not in keep_indices:
            slide_id = prs.slides[i].slide_id
            elem = xml_slides.find(f'.//p:sldId[@id="{slide_id}"]', namespaces=ns)
            if elem is not None:
                xml_slides.remove(elem)
                
    prs.save(file_path)
    print("Slides deleted.")

if __name__ == "__main__":
    delete_slides('course_materials/semester_4/lectures/lecture_02_thermodynamics.pptx')
