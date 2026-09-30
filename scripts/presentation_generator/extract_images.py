import docx, re

doc = docx.Document('course_materials/literature/Электронный учебник по термодинамике.docx')
with open('img_map.txt', 'w', encoding='utf-8') as out:
    for i, p in enumerate(doc.paragraphs):
        xml = p._p.xml
        if 'blip' in xml:
            rids = re.findall(r'r:embed="([^"]+)"', xml)
            for rid in rids:
                if rid in doc.part.rels:
                    tgt = doc.part.rels[rid].target_ref
                    if 1150 <= i <= 1450:
                        pre = ''
                        for off in range(1, 4):
                            if i - off >= 0 and doc.paragraphs[i-off].text.strip():
                                pre = doc.paragraphs[i-off].text.strip()[:100]
                                break
                        out.write(f'Para {i} | {tgt} | {pre}\n')
print('Done mapping!')
